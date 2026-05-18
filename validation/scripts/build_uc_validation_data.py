#!/usr/bin/env python3
"""Split the combined UCSF-supplied UC-system xlsx into per-institution
long-format CSVs that `scripts/run_external_validation.py` can consume.

Input:
    external_validation/DisambiguationEditsForPaul-UCSF.xlsx
        sheet=Added    (InstitutionAbbreviation, PersonID, PMID Manually Added)
        sheet=Removed  (InstitutionAbbreviation, PersonID, PMID Manually Removed)
        sheet=Names    (PersonID, FirstName, MiddleName, LastName, InstitutionAbbreviation)

Output:
    external_validation/uc_system/data/<inst>_data.csv
        columns: PersonID, FirstName, MiddleName, LastName, PMID, Assertion

    where Assertion ∈ {ACCEPTED, REJECTED}; one row per (person, pmid) pair.
    Persons with no Added/Removed PMIDs still get one identity-only row.

Six institutions emitted: ucsf, ucdavis, ucsd, ucla, uci, usc.

Note: Output files contain identifiable researcher data (PersonID + name +
publication assertions). Do NOT commit them. Only the script and the
per-institution row-count summary are safe to commit.
"""
from __future__ import annotations

import csv
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl


REPO_ROOT = Path(__file__).resolve().parents[1]
INPUT_XLSX = REPO_ROOT / 'external_validation' / 'DisambiguationEditsForPaul-UCSF.xlsx'
OUTPUT_DIR = REPO_ROOT / 'external_validation' / 'uc_system' / 'data'

# Map InstitutionAbbreviation (as it appears in the xlsx) to our config short
# names. Order matters for deterministic output.
INSTITUTION_MAP = [
    ('UCSF',     'ucsf'),
    ('UC Davis', 'ucdavis'),
    ('UCSD',     'ucsd'),
    ('UCLA',     'ucla'),
    ('UCI',      'uci'),
    ('USC',      'usc'),
]


def load_names(wb) -> dict:
    """Return {PersonID: (FirstName, MiddleName, LastName, InstitutionAbbr)}."""
    ws = wb['Names']
    names = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        pid, fn, mn, ln, inst = r
        if pid is None:
            continue
        names[int(pid)] = (
            (fn or '').strip(),
            (mn or '').strip(),
            (ln or '').strip(),
            (inst or '').strip(),
        )
    return names


def load_assertions(wb, sheet_name: str, assertion: str) -> dict:
    """Return {InstitutionAbbr: list of (PersonID, PMID, Assertion)}."""
    ws = wb[sheet_name]
    by_inst = defaultdict(list)
    skipped = 0
    for r in ws.iter_rows(min_row=2, values_only=True):
        inst, pid, pmid = r[:3]
        if pid is None or pmid is None or not inst:
            skipped += 1
            continue
        try:
            pid_i = int(pid)
            pmid_i = int(pmid)
        except (TypeError, ValueError):
            skipped += 1
            continue
        by_inst[inst.strip()].append((pid_i, pmid_i, assertion))
    if skipped:
        print(f'  {sheet_name}: skipped {skipped} malformed rows', file=sys.stderr)
    return by_inst


def main():
    if not INPUT_XLSX.exists():
        print(f'ERROR: input not found: {INPUT_XLSX}', file=sys.stderr)
        sys.exit(1)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    print(f'Loading {INPUT_XLSX} ...')
    wb = openpyxl.load_workbook(INPUT_XLSX, read_only=True, data_only=True)

    names = load_names(wb)
    print(f'  Loaded {len(names)} identity rows from Names sheet')

    added = load_assertions(wb, 'Added', 'ACCEPTED')
    total_added = sum(len(v) for v in added.values())
    print(f'  Loaded {total_added} ACCEPTED rows from Added sheet')

    removed = load_assertions(wb, 'Removed', 'REJECTED')
    total_removed = sum(len(v) for v in removed.values())
    print(f'  Loaded {total_removed} REJECTED rows from Removed sheet')

    # Group identities by institution
    inst_to_pids = defaultdict(set)
    for pid, (_, _, _, inst) in names.items():
        if inst:
            inst_to_pids[inst].add(pid)

    summary = []
    for xlsx_inst, short in INSTITUTION_MAP:
        pids = inst_to_pids.get(xlsx_inst, set())
        rows_added = added.get(xlsx_inst, [])
        rows_removed = removed.get(xlsx_inst, [])

        # Orphan check: assertion rows for PersonIDs not in Names sheet for this institution
        valid_pids = pids
        orphans = sum(1 for pid, _, _ in rows_added if pid not in valid_pids) + \
                  sum(1 for pid, _, _ in rows_removed if pid not in valid_pids)

        # PersonIDs that appear in Names but have neither Added nor Removed entries
        pids_with_assertions = {pid for pid, _, _ in rows_added} | {pid for pid, _, _ in rows_removed}
        identity_only_pids = valid_pids - pids_with_assertions

        out_path = OUTPUT_DIR / f'{short}_data.csv'
        with out_path.open('w', newline='') as f:
            w = csv.writer(f)
            w.writerow(['PersonID', 'FirstName', 'MiddleName', 'LastName', 'PMID', 'Assertion'])
            # Identity-only rows for persons with no PMIDs
            for pid in sorted(identity_only_pids):
                fn, mn, ln, _ = names[pid]
                w.writerow([pid, fn, mn, ln, '', ''])
            # Combined assertion rows
            for pid, pmid, assertion in sorted(rows_added) + sorted(rows_removed):
                if pid not in valid_pids:
                    continue
                fn, mn, ln, _ = names[pid]
                w.writerow([pid, fn, mn, ln, pmid, assertion])

        wrote = len(identity_only_pids) + sum(1 for pid, _, _ in rows_added if pid in valid_pids) \
                + sum(1 for pid, _, _ in rows_removed if pid in valid_pids)
        summary.append({
            'institution': xlsx_inst,
            'short': short,
            'persons': len(valid_pids),
            'accepted_pmids': sum(1 for pid, _, _ in rows_added if pid in valid_pids),
            'rejected_pmids': sum(1 for pid, _, _ in rows_removed if pid in valid_pids),
            'identity_only_persons': len(identity_only_pids),
            'orphan_assertion_rows': orphans,
            'output_rows': wrote,
            'output_path': str(out_path.relative_to(REPO_ROOT)),
        })

    print('\nPer-institution output summary:')
    print(f'{"Inst":>10} {"Short":>8} {"Persons":>8} {"Accept":>8} {"Reject":>8} {"IdOnly":>7} {"Orphan":>7} {"Rows":>9}  Path')
    for s in summary:
        print(f'{s["institution"]:>10} {s["short"]:>8} {s["persons"]:>8d} '
              f'{s["accepted_pmids"]:>8d} {s["rejected_pmids"]:>8d} '
              f'{s["identity_only_persons"]:>7d} {s["orphan_assertion_rows"]:>7d} '
              f'{s["output_rows"]:>9d}  {s["output_path"]}')

    return summary


if __name__ == '__main__':
    main()
