#!/usr/bin/env python3
"""Run ReCiter external validation for any institution.

Generic replacement for fred_hutch_run_reciter.py. Supports any CSV/Excel
file format with auto-detected column mapping, flexible name parsing,
email inference, and gold standard extraction.

Three phases:
  1. Load identities + gold standard into ReCiter via API
  2. Run feature-generator for each UID (slow -- hours for large datasets)
  3. Collect scores, evaluate against gold standard, save results

Usage:
  # Step 1: Load data (fast, ~30 seconds)
  python3 scripts/run_external_validation.py --institution fredhutch \\
    --data-file "external_validation/Fred Hutch - Reciter_data_20210101-20240630.xlsx" \\
    --email-domains fredhutch.org,uw.edu,fhcrc.org \\
    --load-only

  # Step 2: Run scoring (slow, run in background)
  python3 scripts/run_external_validation.py --institution fredhutch \\
    --data-file "external_validation/Fred Hutch - Reciter_data_20210101-20240630.xlsx" \\
    --score-only

  # Step 3: Evaluate results (fast)
  python3 scripts/run_external_validation.py --institution fredhutch \\
    --data-file "external_validation/Fred Hutch - Reciter_data_20210101-20240630.xlsx" \\
    --evaluate-only

  # All at once:
  python3 scripts/run_external_validation.py --institution fredhutch \\
    --data-file "external_validation/Fred Hutch - Reciter_data_20210101-20240630.xlsx" \\
    --email-domains fredhutch.org,uw.edu,fhcrc.org

Environment variables:
  RECITER_API_URL   ReCiter API base URL (default: http://localhost:5000)
  RECITER_API_KEY   ReCiter API key (required for load/score/cleanup)
  PUBMED_API_KEY    PubMed API key (optional, for institution setup)

The --base-url and --api-key CLI flags override environment variables.
"""

import argparse
import json
import os
import sys
import threading
import time
from collections import defaultdict
from pathlib import Path

# Add src/ to path for evaluation modules
sys.path.insert(0, str(Path(__file__).parent.parent / 'src'))

DEFAULT_BASE_URL = os.environ.get('RECITER_API_URL', 'http://localhost:5000')

# ---------------------------------------------------------------------------
# Column alias map: normalized_name -> (identity_field, sub_field_or_None)
# ---------------------------------------------------------------------------
COLUMN_ALIASES = {
    # UID (required)
    'uid': ('uid', None),
    'userid': ('uid', None),
    'personid': ('uid', None),
    'employeeid': ('uid', None),
    'cwid': ('uid', None),
    'netid': ('uid', None),

    # Name fields (separate columns)
    'firstname': ('primaryName', 'firstName'),
    'first': ('primaryName', 'firstName'),
    'forename': ('primaryName', 'firstName'),
    'givenname': ('primaryName', 'firstName'),
    'lastname': ('primaryName', 'lastName'),
    'last': ('primaryName', 'lastName'),
    'surname': ('primaryName', 'lastName'),
    'familyname': ('primaryName', 'lastName'),
    'middlename': ('primaryName', 'middleName'),
    'middle': ('primaryName', 'middleName'),
    'middleinitial': ('primaryName', 'middleInitial'),

    # Composite name (needs splitting)
    'trackedauthor': ('_composite_name', 'last_comma_first'),
    'fullname': ('_composite_name', 'first_last'),
    'name': ('_composite_name', 'auto'),

    # PMID + assertion (for gold standard extraction)
    'pmid': ('_gold_standard', 'pmid'),
    'assertion': ('_gold_standard', 'assertion'),
    'assertionstatus': ('_gold_standard', 'assertion'),
    'userassertion': ('_gold_standard', 'assertion'),

    # Email
    'email': ('emails', None),
    'emailaddress': ('emails', None),
    'primaryemail': ('primaryEmail', None),

    # Department / Org Unit
    'department': ('organizationalUnits', 'organizationalUnitLabel'),
    'dept': ('organizationalUnits', 'organizationalUnitLabel'),
    'division': ('organizationalUnits', 'organizationalUnitLabel'),
    'organizationalunit': ('organizationalUnits', 'organizationalUnitLabel'),

    # Degree year
    'degreeyear': ('degreeYear', 'doctoralYear'),
    'doctoralyear': ('degreeYear', 'doctoralYear'),
    'phdyear': ('degreeYear', 'doctoralYear'),
    'bacheloryear': ('degreeYear', 'bachelorYear'),

    # ORCID
    'orcid': ('orcid', None),
    'orcidid': ('orcid', None),

    # Title
    'title': ('title', None),
    'jobtitle': ('title', None),

    # Person type
    'persontype': ('personTypes', None),
    'type': ('personTypes', None),

    # Grants
    'grant': ('grants', None),
    'grants': ('grants', None),
    'grantid': ('grants', None),

    # Institutions
    'institution': ('institutions', None),
    'primaryinstitution': ('primaryInstitution', None),

    # Known relationships
    'relationship': ('knownRelationships', None),
    'relationships': ('knownRelationships', None),
    'knownrelationship': ('knownRelationships', None),
    'knownrelationships': ('knownRelationships', None),
    'collaborator': ('knownRelationships', None),
    'mentor': ('knownRelationships', None),

    # Alternate names
    'maidenname': ('_alternate_name', 'maiden'),
    'alternatename': ('_alternate_name', 'alt'),
    'previousname': ('_alternate_name', 'previous'),

    # Article metadata (not identity fields -- marked as _metadata)
    'articletitle': ('_metadata', 'articleTitle'),
    'dateaddedtoentrez': ('_metadata', 'dateAdded'),
    'datepublicationaddedtoentrez': ('_metadata', 'dateAdded'),
}


# ---------------------------------------------------------------------------
# Column mapping functions
# ---------------------------------------------------------------------------

def normalize_header(header: str) -> str:
    """Normalize a column header for alias lookup.

    Lowercases, strips spaces, underscores, hyphens, and dots.
    """
    return (header.lower()
            .replace(' ', '')
            .replace('_', '')
            .replace('-', '')
            .replace('.', '')
            .strip())


def detect_column_mapping(headers: list) -> tuple:
    """Auto-detect column mapping from file headers.

    Returns:
        (mapping, unmapped) where mapping is {original_header: (field, sub_field)}
        and unmapped is a list of headers that could not be mapped.
    """
    mapping = {}
    unmapped = []

    for header in headers:
        normalized = normalize_header(header)
        if normalized in COLUMN_ALIASES:
            mapping[header] = COLUMN_ALIASES[normalized]
        else:
            unmapped.append(header)

    return mapping, unmapped


def display_mapping(mapping: dict, unmapped: list) -> None:
    """Print human-readable mapping for user review."""
    print('\nDetected column mapping:')
    for header, (field, sub) in sorted(mapping.items()):
        target = f'{field}.{sub}' if sub else field
        print(f'  {header:30s} -> {target}')
    if unmapped:
        print(f'\nUnmapped columns (will be ignored):')
        for h in unmapped:
            print(f'  {h}')


def confirm_mapping(mapping: dict, unmapped: list, interactive: bool) -> dict:
    """In interactive mode, let user review and correct mapping.

    In non-interactive mode, auto-accept the detected mapping.
    """
    display_mapping(mapping, unmapped)
    if not interactive:
        print('\n(Non-interactive mode: accepting auto-detected mapping)')
        return mapping

    print('\nAccept this mapping? (y/n, or enter corrections as "ColumnName=field.subfield")')
    response = input('> ').strip().lower()
    if response == 'y' or response == '':
        return mapping

    # Parse corrections: "ColumnName=field.subfield"
    for correction in response.split(','):
        correction = correction.strip()
        if '=' in correction:
            col, target = correction.split('=', 1)
            col = col.strip()
            parts = target.strip().split('.')
            if len(parts) == 2:
                mapping[col] = (parts[0], parts[1])
            elif len(parts) == 1:
                mapping[col] = (parts[0], None)

    return mapping


# ---------------------------------------------------------------------------
# Name parsing
# ---------------------------------------------------------------------------

def parse_composite_name(value: str, format_hint: str) -> dict:
    """Parse a composite name string into name components.

    Args:
        value: The name string (e.g., "Smith, John David" or "John Smith")
        format_hint: One of "last_comma_first", "first_last", or "auto"

    Returns:
        dict with keys: firstName, lastName, middleName, firstInitial, middleInitial
    """
    first_name = ''
    last_name = ''
    middle_name = ''

    if format_hint == 'auto':
        format_hint = 'last_comma_first' if ',' in value else 'first_last'

    if format_hint == 'last_comma_first':
        parts = value.split(',', 1)
        last_name = parts[0].strip()
        if len(parts) > 1:
            rest = parts[1].strip().split()
            first_name = rest[0] if rest else ''
            middle_name = ' '.join(rest[1:]) if len(rest) > 1 else ''
    elif format_hint == 'first_last':
        parts = value.strip().split()
        if len(parts) == 1:
            first_name = parts[0]
        elif len(parts) == 2:
            first_name = parts[0]
            last_name = parts[1]
        else:
            first_name = parts[0]
            last_name = parts[-1]
            middle_name = ' '.join(parts[1:-1])

    first_initial = first_name[0] if first_name else ''
    middle_initial = middle_name[0] if middle_name else ''

    return {
        'firstName': first_name,
        'lastName': last_name,
        'middleName': middle_name,
        'firstInitial': first_initial,
        'middleInitial': middle_initial,
    }


# ---------------------------------------------------------------------------
# Identity record construction
# ---------------------------------------------------------------------------

def build_identity_record(row: dict, mapping: dict, uid_prefix: str,
                          email_domains: list, institution_label: str) -> dict:
    """Build a ReCiter Identity JSON from a mapped data row.

    Args:
        row: dict of {column_header: value}
        mapping: dict of {column_header: (field, sub_field)}
        uid_prefix: prefix for all UIDs (e.g., "fh_")
        email_domains: list of email domains for inferring emails
        institution_label: institution name for organizationalUnits

    Returns:
        Complete Identity JSON matching ReCiter API contract, or None if
        mandatory fields (uid, firstName, lastName) are missing.
    """
    uid_raw = None
    first_name = ''
    last_name = ''
    middle_name = ''
    middle_initial = ''
    explicit_emails = []
    departments = []
    orcid_value = None
    degree_year = {}
    person_types = []
    grants_list = []
    title_value = None
    institutions_list = []
    known_relationships = []
    alternate_names_extra = []

    # Extract mapped fields
    for header, (field, sub) in mapping.items():
        value = row.get(header)
        if value is None or (isinstance(value, str) and not value.strip()):
            continue
        value = str(value).strip()

        if field == 'uid':
            uid_raw = value
        elif field == 'primaryName':
            if sub == 'firstName':
                first_name = value
            elif sub == 'lastName':
                last_name = value
            elif sub == 'middleName':
                middle_name = value
            elif sub == 'middleInitial':
                middle_initial = value
        elif field == '_composite_name':
            parsed = parse_composite_name(value, sub)
            if not first_name:
                first_name = parsed['firstName']
            if not last_name:
                last_name = parsed['lastName']
            if not middle_name:
                middle_name = parsed['middleName']
            if not middle_initial and parsed['middleInitial']:
                middle_initial = parsed['middleInitial']
        elif field == 'emails':
            explicit_emails.append(value)
        elif field == 'primaryEmail':
            explicit_emails.append(value)
        elif field == 'organizationalUnits':
            departments.append(value)
        elif field == 'orcid':
            orcid_value = value
        elif field == 'degreeYear':
            try:
                degree_year[sub] = int(float(value))
            except (ValueError, TypeError):
                pass
        elif field == 'personTypes':
            person_types.append(value)
        elif field == 'grants':
            grants_list.append(value)
        elif field == 'title':
            title_value = value
        elif field == 'institutions':
            institutions_list.append(value)
        elif field == 'knownRelationships':
            # Parse relationship value -- may be a name string or "Name (type)" format
            rel_value = value
            rel_type = 'colleague'  # default type
            # Check for "Name (type)" pattern
            if '(' in value and value.endswith(')'):
                name_part, type_part = value.rsplit('(', 1)
                rel_value = name_part.strip()
                rel_type = type_part.rstrip(')').strip()
            # Parse the name
            if ',' in rel_value:
                parts = rel_value.split(',', 1)
                rel_last = parts[0].strip()
                rel_first = parts[1].strip().split()[0] if parts[1].strip() else ''
            else:
                name_parts = rel_value.strip().split()
                rel_first = name_parts[0] if name_parts else ''
                rel_last = name_parts[-1] if len(name_parts) > 1 else ''
            known_relationships.append({
                'name': {
                    'firstName': rel_first,
                    'lastName': rel_last,
                },
                'type': rel_type,
            })
        elif field == '_alternate_name':
            alternate_names_extra.append({
                'type': sub,
                'value': value,
            })
        # _gold_standard and _metadata are not identity fields

    # Validate mandatory fields
    if not uid_raw:
        return None
    if not first_name or not last_name:
        return None

    # Compute initials
    uid = f'{uid_prefix}{uid_raw}'
    first_initial = first_name[0] if first_name else ''
    if middle_name and not middle_initial:
        middle_initial = middle_name[0]

    # Build emails: uid_raw@each_domain + explicit emails
    emails = [f'{uid_raw}@{d}' for d in email_domains]
    for email in explicit_emails:
        if email not in emails:
            emails.append(email)

    # Build primary name entry
    primary_name = {
        'firstName': first_name,
        'lastName': last_name,
        'firstInitial': first_initial,
        'middleInitial': middle_initial,
    }

    # Build alternateNames: always include primary, plus any maiden/alternate
    alternate_names = [dict(primary_name)]  # copy

    for alt in alternate_names_extra:
        alt_entry = {
            'firstName': first_name,
            'lastName': alt['value'],
            'firstInitial': first_initial,
            'middleInitial': middle_initial,
        }
        alternate_names.append(alt_entry)

    # Build organizationalUnits
    org_units = [{'organizationalUnitLabel': institution_label}]
    for dept in departments:
        if dept != institution_label:
            org_units.append({'organizationalUnitLabel': dept})

    # Construct record
    record = {
        'uid': uid,
        'primaryName': primary_name,
        'alternateNames': alternate_names,
        'emails': emails,
        'organizationalUnits': org_units,
    }

    # Add optional fields
    if orcid_value:
        record['orcid'] = orcid_value
    if degree_year:
        record['degreeYear'] = degree_year
    if person_types:
        record['personTypes'] = person_types
    if grants_list:
        record['grants'] = grants_list
    if title_value:
        record['title'] = title_value
    if institutions_list:
        record['institutions'] = institutions_list
    if known_relationships:
        record['knownRelationships'] = known_relationships

    return record


# ---------------------------------------------------------------------------
# Gold standard extraction
# ---------------------------------------------------------------------------

def extract_gold_standard(rows: list, mapping: dict, uid_prefix: str) -> dict:
    """Extract gold standard records from data rows.

    Groups rows by UID and collects ACCEPTED/REJECTED PMIDs.
    Non-standard assertions (e.g., "NEEDS REVIEWED") are excluded with a warning.

    Args:
        rows: list of row dicts
        mapping: column mapping dict
        uid_prefix: UID prefix string

    Returns:
        dict of {prefixed_uid: {"knownPmids": [int], "rejectedPmids": [int]}}
    """
    # Find uid, pmid, and assertion columns from mapping
    uid_col = None
    pmid_col = None
    assertion_col = None

    for header, (field, sub) in mapping.items():
        if field == 'uid':
            uid_col = header
        elif field == '_gold_standard' and sub == 'pmid':
            pmid_col = header
        elif field == '_gold_standard' and sub == 'assertion':
            assertion_col = header

    if not uid_col or not pmid_col or not assertion_col:
        return {}

    gs = defaultdict(lambda: {'knownPmids': [], 'rejectedPmids': []})
    excluded_assertions = set()

    for row in rows:
        uid_raw = row.get(uid_col)
        pmid_raw = row.get(pmid_col)
        assertion_raw = row.get(assertion_col)

        if not uid_raw or not pmid_raw or not assertion_raw:
            continue

        uid_raw = str(uid_raw).strip()
        assertion = str(assertion_raw).strip().upper()
        prefixed_uid = f'{uid_prefix}{uid_raw}'

        try:
            pmid = int(float(str(pmid_raw)))
        except (ValueError, TypeError):
            continue

        if assertion == 'ACCEPTED':
            gs[prefixed_uid]['knownPmids'].append(pmid)
        elif assertion == 'REJECTED':
            gs[prefixed_uid]['rejectedPmids'].append(pmid)
        else:
            excluded_assertions.add(assertion)

    if excluded_assertions:
        print(f'  Warning: excluded non-standard assertions: {excluded_assertions}')

    return dict(gs)


# ---------------------------------------------------------------------------
# Data file parsing
# ---------------------------------------------------------------------------

def parse_data_file(file_path: str, sheet: str = None, delimiter: str = None) -> tuple:
    """Parse a CSV or Excel file into headers and row dicts.

    Args:
        file_path: path to the data file
        sheet: Excel sheet name (default: first sheet)
        delimiter: CSV delimiter (default: auto-detect)

    Returns:
        (headers: list[str], rows: list[dict])
    """
    file_path = Path(file_path)
    ext = file_path.suffix.lower()

    if ext in ('.xlsx', '.xls'):
        import openpyxl
        wb = openpyxl.load_workbook(file_path, read_only=True)
        ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]

        row_iter = ws.iter_rows(values_only=True)
        headers = [str(h) if h is not None else '' for h in next(row_iter)]

        rows = []
        for row_values in row_iter:
            row_dict = {}
            for i, val in enumerate(row_values):
                if i < len(headers):
                    row_dict[headers[i]] = val
            rows.append(row_dict)

        wb.close()
        return headers, rows

    elif ext in ('.csv', '.tsv', '.txt'):
        import pandas as pd
        sep = delimiter or (None if ext != '.tsv' else '\t')
        df = pd.read_csv(file_path, sep=sep, engine='python' if sep is None else 'c',
                         dtype=str, keep_default_na=False)
        headers = list(df.columns)
        rows = df.to_dict('records')
        return headers, rows

    else:
        raise ValueError(f'Unsupported file type: {ext}. Use .xlsx, .xls, .csv, or .tsv')


# ---------------------------------------------------------------------------
# API interaction functions (adapted from fred_hutch_run_reciter.py)
# ---------------------------------------------------------------------------

def load_identities(identity_records: list, base_url: str, api_key: str) -> dict:
    """Load identity records into ReCiter via PUT /reciter/save/identities/.

    Sends records in batches of 100.
    """
    import requests

    url = f'{base_url}/reciter/save/identities/'
    headers = {'api-key': api_key, 'Content-Type': 'application/json'}

    batch_size = 100
    results = {'loaded': 0, 'errors': 0}
    for i in range(0, len(identity_records), batch_size):
        batch = identity_records[i:i + batch_size]
        try:
            resp = requests.put(url, json=batch, headers=headers, timeout=60)
            if resp.status_code == 200:
                results['loaded'] += len(batch)
            else:
                print(f'  Batch {i // batch_size + 1}: HTTP {resp.status_code} - {resp.text[:200]}')
                results['errors'] += len(batch)
        except Exception as e:
            print(f'  Batch {i // batch_size + 1}: {e}')
            results['errors'] += len(batch)
        print(f'  Identities: {results["loaded"]}/{len(identity_records)} loaded', end='\r')

    print()
    return results


def load_gold_standard(gs_records: list, base_url: str, api_key: str) -> dict:
    """Load gold standard records into ReCiter via PUT /reciter/goldstandard.

    Sends records in batches of 50.
    """
    import requests

    url = f'{base_url}/reciter/goldstandard'
    headers = {'api-key': api_key, 'Content-Type': 'application/json'}
    params = {'goldStandardUpdateFlag': 'REFRESH'}

    batch_size = 50
    results = {'loaded': 0, 'errors': 0}
    for i in range(0, len(gs_records), batch_size):
        batch = gs_records[i:i + batch_size]
        try:
            resp = requests.put(url, json=batch, headers=headers, params=params, timeout=60)
            if resp.status_code == 200:
                results['loaded'] += len(batch)
            else:
                print(f'  GS Batch {i // batch_size + 1}: HTTP {resp.status_code} - {resp.text[:200]}')
                results['errors'] += len(batch)
        except Exception as e:
            print(f'  GS Batch {i // batch_size + 1}: {e}')
            results['errors'] += len(batch)
        print(f'  Gold standard: {results["loaded"]}/{len(gs_records)} loaded', end='\r')

    print()
    return results


# ---------------------------------------------------------------------------
# Adaptive Worker Pool (D-11)
# ---------------------------------------------------------------------------

class AdaptiveWorkerPool:
    """Dynamically adjusts worker count based on response time health.

    Monitors a rolling window of response times. Once a baseline is
    established (after window_size responses), it scales workers down
    when recent averages exceed scale_down_factor * baseline, and
    scales back up when averages drop below scale_up_factor * baseline.
    """

    def __init__(self, initial_workers=4, min_workers=1, max_workers=6,
                 scale_down_factor=2.0, scale_up_factor=1.2, window_size=10):
        self.current_workers = initial_workers
        self.min_workers = min_workers
        self.max_workers = max_workers
        self.scale_down_factor = scale_down_factor
        self.scale_up_factor = scale_up_factor
        self.response_times = []  # rolling window
        self.window_size = window_size
        self.baseline = None  # set after first window_size responses
        self._lock = threading.Lock()

    def record_response_time(self, seconds: float):
        """Record a response time observation."""
        with self._lock:
            self.response_times.append(seconds)
            if len(self.response_times) > self.window_size * 3:
                self.response_times = self.response_times[-self.window_size * 3:]
            if self.baseline is None and len(self.response_times) >= self.window_size:
                self.baseline = sum(self.response_times[:self.window_size]) / self.window_size

    def should_scale_down(self) -> bool:
        """Check if recent response times warrant scaling down."""
        if self.baseline is None or len(self.response_times) < self.window_size:
            return False
        recent_avg = sum(self.response_times[-self.window_size:]) / self.window_size
        return recent_avg > self.baseline * self.scale_down_factor

    def should_scale_up(self) -> bool:
        """Check if recent response times warrant scaling up."""
        if self.baseline is None or len(self.response_times) < self.window_size:
            return False
        recent_avg = sum(self.response_times[-self.window_size:]) / self.window_size
        return recent_avg < self.baseline * self.scale_up_factor

    def adjust(self) -> int:
        """Evaluate and adjust the current worker count. Returns new count."""
        with self._lock:
            if self.should_scale_down() and self.current_workers > self.min_workers:
                self.current_workers -= 1
            elif self.should_scale_up() and self.current_workers < self.max_workers:
                self.current_workers += 1
            return self.current_workers


def _score_one_uid(uid, base_url, api_key, use_gold_standard,
                   filter_by_feedback='ACCEPTED_AND_REJECTED'):
    """Score a single UID. Returns (uid, scores_list) or (uid, None, error_info) on failure."""
    import requests

    url = f'{base_url}/reciter/feature-generator/by/uid'
    headers = {'api-key': api_key}
    params = {
        'uid': uid,
        'useGoldStandard': use_gold_standard,
        'filterByFeedback': filter_by_feedback,
        'analysisRefreshFlag': 'true',
        'retrievalRefreshFlag': 'FALSE',
    }
    per_uid_timeout = int(os.environ.get('RECITER_PER_UID_TIMEOUT', '300'))
    try:
        resp = requests.get(url, headers=headers, params=params, timeout=per_uid_timeout)
        if resp.status_code == 200:
            data = resp.json()
            articles = data.get('reCiterArticleFeatures', []) if isinstance(data, dict) else data
            uid_scores = []
            for article in articles:
                pmid = article.get('pmid')
                score = article.get('authorshipLikelihoodScore')
                assertion = article.get('userAssertion')
                if pmid is not None and score is not None:
                    uid_scores.append({'pmid': pmid, 'score': score, 'userAssertion': assertion})
            return uid, uid_scores
        else:
            return uid, None, {'status': resp.status_code, 'error': resp.text[:200]}
    except Exception as e:
        return uid, None, {'error': str(e)}


def run_feature_generator(uids: list, base_url: str, api_key: str,
                          scores_file: Path, use_gold_standard: str = 'AS_EVIDENCE',
                          workers: int = 4, adaptive: bool = False,
                          filter_by_feedback: str = 'ACCEPTED_AND_REJECTED') -> dict:
    """Run feature-generator for each UID with concurrent workers.

    Saves progress incrementally to scores_file (resume-safe).

    Args:
        uids: list of UIDs to score
        base_url: ReCiter API base URL
        api_key: ReCiter API key
        scores_file: Path to save/resume scores JSON
        use_gold_standard: 'AS_EVIDENCE' or 'FOR_TESTING_ONLY'
        workers: number of concurrent workers (initial count if adaptive=True)
        adaptive: enable adaptive worker scaling based on response times
    """
    from concurrent.futures import ThreadPoolExecutor, as_completed

    # Load existing scores for resume
    existing_scores = {}
    if scores_file.exists():
        with open(scores_file) as f:
            existing_scores = json.load(f)

    # Filter to unscored UIDs
    to_score = [uid for uid in uids if uid not in existing_scores]
    skipped = len(uids) - len(to_score)
    total = len(uids)

    results = {'completed': 0, 'failed': 0, 'skipped': skipped, 'errors': [],
               'stopped_early': False}
    mode_label = 'adaptive' if adaptive else 'fixed'
    print(f'  Workers: {workers} ({mode_label})')
    print(f'  To score: {len(to_score)}, skipping: {skipped} already done')

    save_lock = threading.Lock()

    # Initialize adaptive pool if enabled
    pool_adapter = AdaptiveWorkerPool(
        initial_workers=workers, min_workers=1, max_workers=max(workers + 2, 6)
    ) if adaptive else None

    def _save_scores():
        """Persist scores to disk."""
        with save_lock:
            with open(scores_file, 'w') as f:
                json.dump(existing_scores, f, indent=2)

    def _process_result(result):
        """Process a single scoring result. Returns (uid, success)."""
        uid = result[0]
        if len(result) == 2 and result[1] is not None:
            with save_lock:
                existing_scores[uid] = result[1]
            return uid, True
        else:
            error_info = result[2] if len(result) > 2 else {'error': 'unknown'}
            error_info['uid'] = uid
            return uid, False

    def _check_failure_threshold(completed, failed, total_to_score):
        """Check if cumulative failure rate exceeds 5% threshold (D-13).

        Requires at least 20 UIDs processed to avoid false triggers on small runs.
        For the full 874-UID run this means >43 failures; for smaller test sets
        it triggers proportionally.
        """
        if total_to_score == 0:
            return False
        processed = completed + failed
        if processed < 20:
            return False
        threshold = float(os.environ.get('RECITER_MAX_FAIL_PCT', '0.05'))
        if failed / total_to_score > threshold:
            return True
        return False

    # --- Main scoring pass ---
    if adaptive and pool_adapter:
        # Batch-based adaptive scoring: process UIDs in batches sized to current_workers
        remaining = list(to_score)
        done_count = 0
        while remaining:
            batch_size = pool_adapter.current_workers
            batch = remaining[:batch_size]
            remaining = remaining[batch_size:]

            with ThreadPoolExecutor(max_workers=batch_size) as executor:
                futures = {
                    executor.submit(_score_one_uid, uid, base_url, api_key,
                                    use_gold_standard, filter_by_feedback): uid
                    for uid in batch
                }
                for future in as_completed(futures):
                    start_t = time.monotonic()
                    result = future.result()
                    elapsed = time.monotonic() - start_t
                    pool_adapter.record_response_time(elapsed)

                    uid, success = _process_result(result)
                    done_count += 1
                    if success:
                        results['completed'] += 1
                    else:
                        results['failed'] += 1
                        error_info = result[2] if len(result) > 2 else {'error': 'unknown'}
                        error_info['uid'] = uid
                        results['errors'].append(error_info)

            # Adjust workers after each batch
            pool_adapter.adjust()

            # Save periodically
            if done_count % 10 == 0 or not remaining:
                _save_scores()
                total_done = results['completed'] + results['skipped']
                pct = total_done / total * 100 if total > 0 else 0
                print(f'  [{total_done}/{total}] {pct:.1f}% done | workers: {pool_adapter.current_workers} | '
                      f'{results["completed"]} scored, {results["failed"]} failed')

            # Check cumulative failure threshold (D-13)
            if _check_failure_threshold(results['completed'], results['failed'], len(to_score)):
                print(f'\n  STOPPING: Cumulative failure rate '
                      f'{results["failed"]}/{len(to_score)} '
                      f'({results["failed"]/len(to_score)*100:.1f}%) exceeds 5% threshold')
                results['stopped_early'] = True
                _save_scores()
                return results
    else:
        # Fixed-worker scoring (original behavior)
        done_count = 0
        with ThreadPoolExecutor(max_workers=workers) as executor:
            futures = {
                executor.submit(_score_one_uid, uid, base_url, api_key,
                                use_gold_standard, filter_by_feedback): uid
                for uid in to_score
            }
            for future in as_completed(futures):
                result = future.result()
                uid, success = _process_result(result)
                done_count += 1
                if success:
                    results['completed'] += 1
                else:
                    results['failed'] += 1
                    error_info = result[2] if len(result) > 2 else {'error': 'unknown'}
                    error_info['uid'] = uid
                    results['errors'].append(error_info)

                # Save every 10 completions
                if done_count % 10 == 0 or done_count == len(to_score):
                    _save_scores()
                    total_done = results['completed'] + results['skipped']
                    pct = total_done / total * 100 if total > 0 else 0
                    print(f'  [{total_done}/{total}] {pct:.1f}% done | {results["completed"]} scored, '
                          f'{results["skipped"]} skipped, {results["failed"]} failed')

                # Check cumulative failure threshold (D-13)
                if _check_failure_threshold(results['completed'], results['failed'], len(to_score)):
                    print(f'\n  STOPPING: Cumulative failure rate '
                          f'{results["failed"]}/{len(to_score)} '
                          f'({results["failed"]/len(to_score)*100:.1f}%) exceeds 5% threshold')
                    results['stopped_early'] = True
                    _save_scores()
                    return results

    _save_scores()

    # --- Retry-at-end pass (D-12) ---
    if results['errors'] and not results['stopped_early']:
        failed_uids = [e['uid'] for e in results['errors'] if 'uid' in e]
        if failed_uids:
            print(f'\n  Retry pass: {len(failed_uids)} failed UIDs (single-threaded)')
            retry_successes = 0
            retry_still_failed = []
            for uid in failed_uids:
                result = _score_one_uid(uid, base_url, api_key, use_gold_standard,
                                        filter_by_feedback)
                _, success = _process_result(result)
                if success:
                    retry_successes += 1
                    results['completed'] += 1
                    results['failed'] -= 1
                else:
                    error_info = result[2] if len(result) > 2 else {'error': 'unknown'}
                    error_info['uid'] = uid
                    retry_still_failed.append(error_info)

            # Replace errors list with only the still-failing UIDs
            results['errors'] = retry_still_failed
            print(f'  Retry results: {retry_successes} recovered, '
                  f'{len(retry_still_failed)} still failed')
            _save_scores()

    # Check failure threshold after retry
    if _check_failure_threshold(results['completed'], results['failed'], len(to_score)):
        results['stopped_early'] = True

    # Final save
    _save_scores()

    return results


def evaluate_scores(scores_file: Path, gold_standard: dict) -> dict:
    """Evaluate collected scores against gold standard."""
    import numpy as np
    from evaluation import evaluate_predictions
    from calibration import compute_ece, verify_extreme_calibration

    with open(scores_file) as f:
        all_scores = json.load(f)

    # Build arrays: label (1=accepted, 0=rejected) and predicted score
    labels = []
    predictions = []
    uid_stats = {'scored': 0, 'missing': 0, 'articles_scored': 0, 'articles_missing_score': 0}

    for uid, gs_data in gold_standard.items():
        if uid not in all_scores:
            uid_stats['missing'] += 1
            continue
        uid_stats['scored'] += 1

        scored_articles = {a['pmid']: a['score'] for a in all_scores[uid]}
        gs_pmids = {**{p: 1 for p in gs_data['knownPmids']}, **{p: 0 for p in gs_data['rejectedPmids']}}

        for pmid, label in gs_pmids.items():
            if pmid in scored_articles:
                labels.append(label)
                predictions.append(scored_articles[pmid] / 100.0)
                uid_stats['articles_scored'] += 1
            else:
                uid_stats['articles_missing_score'] += 1

    labels = np.array(labels)
    predictions = np.array(predictions)

    print(f'\nEvaluation dataset:')
    print(f'  UIDs scored: {uid_stats["scored"]}, missing: {uid_stats["missing"]}')
    print(f'  Articles scored: {uid_stats["articles_scored"]}, missing score: {uid_stats["articles_missing_score"]}')
    print(f'  Labels: {int(labels.sum())} accepted, {int(len(labels) - labels.sum())} rejected')

    # Run evaluation
    eval_results = evaluate_predictions(labels, predictions, tag='external_validation')
    extreme_cal = verify_extreme_calibration(labels, predictions)

    # Review burden at standard thresholds
    scores_100 = predictions * 100
    auto_accept = scores_100 >= 95
    auto_reject = scores_100 <= 10
    needs_review = ~auto_accept & ~auto_reject
    review_burden = {
        'auto_accept_count': int(auto_accept.sum()),
        'auto_accept_accuracy': float(labels[auto_accept].mean()) if auto_accept.any() else None,
        'needs_review_count': int(needs_review.sum()),
        'needs_review_pct': float(needs_review.mean() * 100),
        'auto_reject_count': int(auto_reject.sum()),
        'auto_reject_accuracy': float((1 - labels[auto_reject]).mean()) if auto_reject.any() else None,
    }

    results = {
        'uid_stats': uid_stats,
        'metrics': eval_results,
        'extreme_calibration': extreme_cal,
        'review_burden': review_burden,
    }

    # Print summary
    print(f'\n=== External Validation Results ===')
    print(f'AUC-ROC:        {eval_results["auc_roc"]:.4f}')
    print(f'ECE:            {eval_results["ece"]:.4f}')
    print(f'Brier:          {eval_results["brier"]:.4f}')
    print(f'Review burden:  {review_burden["needs_review_pct"]:.1f}%')
    if review_burden["auto_accept_accuracy"] is not None:
        print(f'Auto-accept accuracy: {review_burden["auto_accept_accuracy"]:.4f}')
    if review_burden["auto_reject_accuracy"] is not None:
        print(f'Auto-reject accuracy: {review_burden["auto_reject_accuracy"]:.4f}')

    return results


def verify_api_connection(base_url: str, api_key: str) -> bool:
    """Pre-run safety check: verify ReCiter API is reachable."""
    import requests
    try:
        resp = requests.get(f'{base_url}/reciter/ping',
                           headers={'api-key': api_key}, timeout=10)
        if resp.status_code == 200 and 'Healthy' in resp.text:
            return True
        print(f'API responded but unhealthy: HTTP {resp.status_code}')
        return False
    except requests.ConnectionError:
        print(f'Cannot connect to ReCiter API at {base_url}')
        print(f'Set RECITER_API_URL environment variable or use --base-url flag.')
        return False
    except requests.Timeout:
        print(f'ReCiter API at {base_url} timed out after 10 seconds')
        return False


def cleanup_via_api(uids: list, base_url: str, api_key: str) -> dict:
    """Delete all prefixed records via ReCiter API DELETE endpoints.

    Deletion order: AnalysisOutput, ESearchResult, GoldStandard, Identity
    (least dependent first). Idempotent -- deleting non-existent records
    returns 404 which is treated as success.
    """
    import requests
    headers = {'api-key': api_key}
    results = {}

    # Order matters: delete dependents before parents
    resources = [
        ('analysis', 'AnalysisOutput'),
        ('esearchresult', 'ESearchResult'),
        ('goldstandard', 'GoldStandard'),
        ('identity', 'Identity'),
    ]

    for resource_path, display_name in resources:
        deleted = 0
        errors = 0
        for uid in uids:
            try:
                resp = requests.delete(
                    f'{base_url}/reciter/{resource_path}/{uid}',
                    headers=headers, timeout=30
                )
                if resp.status_code in (200, 404):
                    deleted += 1
                else:
                    errors += 1
            except Exception as e:
                errors += 1
        results[resource_path] = {'deleted': deleted, 'errors': errors}
        print(f'  {display_name}: {deleted}/{len(uids)} deleted')

    print('Cleanup complete.')
    return results


# ---------------------------------------------------------------------------
# CLI entry point
# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(
        description='Run ReCiter external validation for any institution',
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Examples:
  # Load identities + gold standard
  python3 scripts/run_external_validation.py --institution fredhutch \\
    --data-file data.xlsx --email-domains fredhutch.org,uw.edu --load-only

  # Run scoring (resume-safe)
  python3 scripts/run_external_validation.py --institution fredhutch \\
    --data-file data.xlsx --score-only

  # Evaluate results
  python3 scripts/run_external_validation.py --institution fredhutch \\
    --data-file data.xlsx --evaluate-only
""")

    parser.add_argument('--institution', required=True,
                        help='Institution short name (e.g., fredhutch, ucsf)')
    parser.add_argument('--data-file', required=True,
                        help='Path to CSV/Excel data file with identity + gold standard data')
    parser.add_argument('--config',
                        help='Path to config JSON from institution setup script')
    parser.add_argument('--uid-prefix',
                        help='UID prefix (default: first 2 chars of institution name + _)')
    parser.add_argument('--email-domains',
                        help='Comma-separated email domains for identity records')
    parser.add_argument('--institution-label',
                        help='Full institution name for organizationalUnits')
    parser.add_argument('--sheet',
                        help='Excel sheet name (default: first sheet)')
    parser.add_argument('--delimiter',
                        help='CSV delimiter (default: auto-detect)')
    parser.add_argument('--non-interactive', action='store_true',
                        help='Accept auto-detected column mapping without confirmation')
    parser.add_argument('--load-only', action='store_true',
                        help='Only load identities + gold standard')
    parser.add_argument('--score-only', action='store_true',
                        help='Only run feature-generator (resume-safe)')
    parser.add_argument('--evaluate-only', action='store_true',
                        help='Only evaluate collected scores')
    parser.add_argument('--cleanup', action='store_true',
                        help='Delete all prefixed records via API (alias for --revert)')
    parser.add_argument('--revert', action='store_true',
                        help='Delete all prefixed records via API (cleanup after validation run)')
    parser.add_argument('--api-key',
                        help='ReCiter API key (overrides RECITER_API_KEY env var)')
    parser.add_argument('--dry-run', action='store_true',
                        help='With --revert: show what would be deleted without deleting')
    parser.add_argument('--base-url', default=DEFAULT_BASE_URL,
                        help=f'ReCiter base URL (default: {DEFAULT_BASE_URL})')
    parser.add_argument('--workers', type=int, default=4,
                        help='Concurrent scoring workers (default: 4, initial count if --adaptive)')
    parser.add_argument('--adaptive', action='store_true',
                        help='Enable adaptive worker scaling based on response times (D-11)')
    parser.add_argument('--use-gold-standard', default='AS_EVIDENCE',
                        choices=['AS_EVIDENCE', 'FOR_TESTING_ONLY'],
                        help='How to use gold standard during scoring')
    parser.add_argument('--sample',
                        help='JSON file with UID subset to process')
    parser.add_argument('--filter-by-feedback', default='ACCEPTED_AND_REJECTED',
                        choices=['ALL', 'ACCEPTED_ONLY', 'REJECTED_ONLY',
                                 'ACCEPTED_AND_NULL', 'REJECTED_AND_NULL',
                                 'ACCEPTED_AND_REJECTED', 'NULL'],
                        help='filterByFeedback passed to feature-generator. '
                             'ALL includes PENDING (new-match) articles; '
                             'scores then write to <inst>_scores_all.json.')

    args = parser.parse_args()

    # Derive defaults
    uid_prefix = args.uid_prefix or (args.institution[:2] + '_')
    email_domains = args.email_domains.split(',') if args.email_domains else []
    institution_label = args.institution_label or args.institution
    interactive = not args.non_interactive

    # Load config if provided
    config = {}
    if args.config:
        with open(args.config) as f:
            config = json.load(f)
        # Extract email domains from config if not provided via CLI
        if not email_domains and 'strategy.email.default.suffixes' in config:
            suffixes = config['strategy.email.default.suffixes']
            email_domains = [s.strip().lstrip('@') for s in suffixes.split(',')]
        if institution_label == args.institution and 'strategy.authorAffiliationScoringStrategy.homeInstitution-label' in config:
            institution_label = config['strategy.authorAffiliationScoringStrategy.homeInstitution-label']

    # Connection config: CLI flag > env var > default
    api_key = args.api_key or os.environ.get('RECITER_API_KEY')
    base_url = args.base_url  # argparse default already reads from RECITER_API_URL via DEFAULT_BASE_URL

    # Results directory
    results_dir = Path(__file__).parent.parent / 'external_validation' / 'results' / args.institution
    run_suffix = '_all' if args.filter_by_feedback == 'ALL' else ''
    scores_file = results_dir / f'{args.institution}_scores{run_suffix}.json'

    if args.cleanup or args.revert:
        if not api_key:
            print('Error: API key required for cleanup. Set RECITER_API_KEY or use --api-key.')
            sys.exit(1)

        print(f'{"[DRY RUN] " if args.dry_run else ""}Cleaning up all {uid_prefix}* records via API...')
        headers_data, rows = parse_data_file(args.data_file, sheet=args.sheet, delimiter=args.delimiter)
        mapping, _ = detect_column_mapping(headers_data)
        uid_col = None
        for header, (field, sub) in mapping.items():
            if field == 'uid':
                uid_col = header
                break
        if uid_col:
            uids = sorted(set(f'{uid_prefix}{str(row[uid_col]).strip()}' for row in rows if row.get(uid_col)))
            print(f'  Found {len(uids)} prefixed UIDs to clean up')
            if args.dry_run:
                print(f'  Would delete from: AnalysisOutput, ESearchResult, GoldStandard, Identity')
                for uid in uids[:10]:
                    print(f'    {uid}')
                if len(uids) > 10:
                    print(f'    ... and {len(uids) - 10} more')
            else:
                cleanup_via_api(uids, base_url, api_key)
        return

    run_all = not args.load_only and not args.score_only and not args.evaluate_only

    if (run_all or args.load_only or args.score_only) and not api_key:
        print('Error: ReCiter API key required.')
        print('  Set RECITER_API_KEY environment variable or use --api-key flag.')
        sys.exit(1)

    results_dir.mkdir(parents=True, exist_ok=True)

    # Pre-run safety checks
    if run_all or args.load_only or args.score_only:
        if not uid_prefix or uid_prefix.strip() == '':
            print('Error: UID prefix is required. Use --uid-prefix or ensure --institution is set.')
            sys.exit(1)

        if not verify_api_connection(base_url, api_key):
            sys.exit(1)
        print(f'  API connection verified: {base_url}')

    # Parse data file
    print(f'Parsing {args.data_file}...')
    headers, rows = parse_data_file(args.data_file, sheet=args.sheet, delimiter=args.delimiter)
    print(f'  {len(rows)} rows, {len(headers)} columns')

    # Detect column mapping
    mapping, unmapped = detect_column_mapping(headers)
    mapping = confirm_mapping(mapping, unmapped, interactive)

    # Build identity records from distinct UIDs
    print('\nBuilding identity records...')
    identities = {}
    skipped = 0
    for row in rows:
        record = build_identity_record(row, mapping, uid_prefix, email_domains, institution_label)
        if record and record['uid'] not in identities:
            identities[record['uid']] = record
        elif not record:
            skipped += 1
    print(f'  {len(identities)} unique identities ({skipped} rows skipped)')

    # Extract gold standard
    gs = extract_gold_standard(rows, mapping, uid_prefix)
    print(f'  {len(gs)} gold standard records')

    uids = sorted(identities.keys())

    if len(uids) > 500:
        print(f'\n  WARNING: Large run with {len(uids)} UIDs. This may take many hours.')
        print(f'  Consider using --sample to test with a subset first.')

    # Filter to sample if provided
    if args.sample:
        with open(args.sample) as f:
            sample_raw = json.load(f)
        sample_uids = set(f'{uid_prefix}{u}' if not u.startswith(uid_prefix) else u for u in sample_raw)
        identities = {k: v for k, v in identities.items() if k in sample_uids}
        gs = {k: v for k, v in gs.items() if k in sample_uids}
        uids = sorted(identities.keys())
        print(f'  Sample mode: {len(uids)} UIDs from {args.sample}')

    # Phase 1: Load
    if run_all or args.load_only:
        print(f'\n--- Phase 1: Loading {len(identities)} identities ---')
        id_results = load_identities(list(identities.values()), args.base_url, api_key)
        print(f'  Done: {id_results["loaded"]} loaded, {id_results["errors"]} errors')

        gs_records = [{'uid': uid, **data} for uid, data in gs.items()]
        print(f'\n--- Phase 1: Loading {len(gs_records)} gold standard records ---')
        gs_results = load_gold_standard(gs_records, args.base_url, api_key)
        print(f'  Done: {gs_results["loaded"]} loaded, {gs_results["errors"]} errors')

    # Phase 2: Score
    if run_all or args.score_only:
        print(f'\n--- Phase 2: Running feature-generator ({len(uids)} UIDs) ---')
        print(f'  Mode: useGoldStandard={args.use_gold_standard}, '
              f'filterByFeedback={args.filter_by_feedback}')
        print(f'  Scores file: {scores_file}')
        print(f'  Resume-safe: existing scores will be skipped')
        score_results = run_feature_generator(uids, args.base_url, api_key,
                                              scores_file, args.use_gold_standard,
                                              workers=args.workers,
                                              adaptive=args.adaptive,
                                              filter_by_feedback=args.filter_by_feedback)
        print(f'\n  Done: {score_results["completed"]} scored, {score_results["skipped"]} skipped, '
              f'{score_results["failed"]} failed')
        if score_results['errors']:
            errors_file = results_dir / f'{args.institution}_scoring_errors{run_suffix}.json'
            with open(errors_file, 'w') as f:
                json.dump(score_results['errors'], f, indent=2)
            print(f'  Errors saved to: {errors_file}')

    # Phase 3: Evaluate
    if run_all or args.evaluate_only:
        if not scores_file.exists():
            print(f'\nNo scores file found at {scores_file}. Run --score-only first.')
            sys.exit(1)

        print(f'\n--- Phase 3: Evaluating scores ---')
        eval_results = evaluate_scores(scores_file, gs)

        # Save results
        eval_file = results_dir / f'{args.institution}_evaluation_results.json'
        with open(eval_file, 'w') as f:
            json.dump(eval_results, f, indent=2, default=str)
        print(f'\nResults saved to: {eval_file}')


if __name__ == '__main__':
    main()
